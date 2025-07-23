from docx import Document
from docx.shared import Pt


def highlight_letters_with_font_change(input_file, output_file, target_word):

    doc = Document(input_file)

    target_letters = list(target_word.lower())
    letter_index = 0


    for paragraph in doc.paragraphs:
        new_runs = []
        for run in paragraph.runs:
            new_text = []

            for char in run.text:
                if letter_index < len(target_letters) and char.lower() == target_letters[letter_index]:
                    highlighted_run = (char, True)
                    new_runs.append(highlighted_run)
                    letter_index += 1
                else:
                    regular_run = (char, False)
                    new_runs.append(regular_run)

            if letter_index >= len(target_letters):
                letter_index = 0

        paragraph.clear()

        for char, is_highlighted in new_runs:
            run = paragraph.add_run(char)
            if is_highlighted:
                run.font.name = "Arial black"

    doc.save(output_file)
    print(f"Файл успешно сохранён: {output_file}")


if __name__ == "__main__":
    input_file = "input.docx"  # Входной файл Word
    output_file = "output.docx"  # Выходной файл Word
    target_word = "Секрет"  # Целевое слово

    highlight_letters_with_font_change(input_file, output_file, target_word)
